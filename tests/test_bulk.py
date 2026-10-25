"""
Tests for the bulk audit feature — column mapping, row-by-row comparison,
and the bulk exporters. All offline, no network needed.
"""

from core import (
    bulk_overview,
    guess_column_mapping,
    guess_identifier_column,
    run_bulk_audit,
)
from exports import (
    _build_file_centric_bulk_table,
    to_csv_bytes_bulk,
    to_excel_bytes_bulk,
    to_pdf_bytes_bulk,
)

LUBE_SCHEMA = {
    "template_name": "Lube Oil List Audit",
    "fields": [
        {"field_name": "recomm_lubricant", "label": "Recommended Lubricant", "required": True,
         "acceptable_values": ["Shell Omala S2 G220", "220# Gear Oil"]},
        {"field_name": "recomm_grade", "label": "Recommended Grade", "required": True,
         "acceptable_values": ["ISO VG220", "ISO VG32"]},
        {"field_name": "uom", "label": "UOM", "required": True, "acceptable_values": ["Litres"]},
    ],
}

REAL_COLUMNS = ["SI No", "Equipment Type", "Equipment Tag", "Equipment Description",
                "Area to be lubricated", "Volume Ltrs", "UOM", "Recomm. Lubricant", "Recomm. Grade"]


# ---------- guess_column_mapping ----------

def test_guess_column_mapping_exact_and_fuzzy():
    mapping = guess_column_mapping(LUBE_SCHEMA, REAL_COLUMNS)
    assert mapping["uom"] == "UOM"
    assert mapping["recomm_lubricant"] == "Recomm. Lubricant"
    assert mapping["recomm_grade"] == "Recomm. Grade"


def test_guess_column_mapping_no_match_returns_none():
    mapping = guess_column_mapping(LUBE_SCHEMA, ["Totally Unrelated Column"])
    # a wildly unrelated single column shouldn't be force-matched to everything
    matched_count = sum(1 for v in mapping.values() if v is not None)
    assert matched_count <= 1


# ---------- guess_identifier_column ----------

def test_guess_identifier_column_prefers_equipment_tag():
    assert guess_identifier_column(REAL_COLUMNS) == "Equipment Tag"


def test_guess_identifier_column_falls_back_to_first():
    assert guess_identifier_column(["Foo", "Bar"]) == "Foo"


def test_guess_identifier_column_empty_list():
    assert guess_identifier_column([]) is None


# ---------- run_bulk_audit ----------

def _make_records():
    return [
        {"Equipment Tag": "MA-EPE255-F000643A/B", "UOM": "Litres",
         "Recomm. Lubricant": "Shell Omala S2 G220", "Recomm. Grade": "ISO VG220"},
        {"Equipment Tag": "MF-EPE255-A112101", "UOM": "Litres",
         "Recomm. Lubricant": None, "Recomm. Grade": "ISO VG32"},
        {"Equipment Tag": "MP-EPE255-P125003", "UOM": "Gallons",
         "Recomm. Lubricant": "Unbranded Oil", "Recomm. Grade": "ISO VG220"},
    ]


def test_run_bulk_audit_row_count_matches_input():
    records = _make_records()
    mapping = guess_column_mapping(LUBE_SCHEMA, list(records[0].keys()))
    results = run_bulk_audit(records, LUBE_SCHEMA, mapping, "Equipment Tag")
    assert len(results) == 3


def test_run_bulk_audit_uses_identifier_column():
    records = _make_records()
    mapping = guess_column_mapping(LUBE_SCHEMA, list(records[0].keys()))
    results = run_bulk_audit(records, LUBE_SCHEMA, mapping, "Equipment Tag")
    assert results[0]["identifier"] == "MA-EPE255-F000643A/B"


def test_run_bulk_audit_falls_back_to_row_number_when_identifier_blank():
    records = [{"Equipment Tag": "", "UOM": "Litres"}]
    mapping = {"recomm_lubricant": None, "recomm_grade": None, "uom": "UOM"}
    results = run_bulk_audit(records, LUBE_SCHEMA, mapping, "Equipment Tag")
    assert results[0]["identifier"] == "Row 1"


def test_run_bulk_audit_detects_missing_and_mismatched():
    records = _make_records()
    mapping = guess_column_mapping(LUBE_SCHEMA, list(records[0].keys()))
    results = run_bulk_audit(records, LUBE_SCHEMA, mapping, "Equipment Tag")

    # Row 1: everything valid
    assert results[0]["summary"]["MATCHED"] == 3

    # Row 2: missing lubricant
    assert results[1]["summary"]["MISSING"] == 1

    # Row 3: wrong UOM and wrong lubricant
    assert results[2]["summary"]["MISMATCHED"] == 2


def test_run_bulk_audit_unmapped_field_is_always_missing():
    records = [{"Equipment Tag": "X"}]
    mapping = {"recomm_lubricant": None, "recomm_grade": None, "uom": None}
    results = run_bulk_audit(records, LUBE_SCHEMA, mapping, "Equipment Tag")
    assert results[0]["summary"]["MISSING"] == 3


# ---------- bulk_overview ----------

def test_bulk_overview_flags_partial_matches():
    records = _make_records()
    mapping = guess_column_mapping(LUBE_SCHEMA, list(records[0].keys()))
    results = run_bulk_audit(records, LUBE_SCHEMA, mapping, "Equipment Tag")
    overview = bulk_overview(results)

    assert overview[0]["all_matched"] is True
    assert overview[1]["all_matched"] is False
    assert overview[2]["all_matched"] is False


# ---------- bulk exporters ----------

def _sample_results():
    records = _make_records()
    mapping = guess_column_mapping(LUBE_SCHEMA, list(records[0].keys()))
    return run_bulk_audit(records, LUBE_SCHEMA, mapping, "Equipment Tag")


def test_bulk_excel_export_valid():
    data = to_excel_bytes_bulk(_sample_results(), "Lube Oil List Audit")
    assert data[:2] == b"PK"
    assert len(data) > 3000


def test_bulk_file_centric_table_has_one_row_per_file_and_field_columns():
    headers, rows, statuses = _build_file_centric_bulk_table(_sample_results(), "File")

    assert headers == [
        "File",
        "Recommended Lubricant",
        "Recommended Grade",
        "UOM",
        "Overall Status",
        "Review Fields",
    ]
    assert len(rows) == 3
    assert rows[0][0] == "MA-EPE255-F000643A/B"
    assert rows[0][1:4] == ["Shell Omala S2 G220", "ISO VG220", "Litres"]
    assert rows[0][4] == "MATCHED"
    assert rows[1][4] == "REVIEW"
    assert "Recommended Lubricant" in rows[1][5]
    assert statuses[2]["UOM"] == "MISMATCHED"


def test_bulk_csv_export_starts_with_file_centric_table():
    text = to_csv_bytes_bulk(_sample_results(), "Lube Oil List Audit", "File").decode("utf-8")
    lines = text.splitlines()

    assert lines[3] == "File,Recommended Lubricant,Recommended Grade,UOM,Overall Status,Review Fields"
    assert lines[4].startswith("MA-EPE255-F000643A/B,Shell Omala S2 G220,ISO VG220,Litres,MATCHED,")
    assert "Detailed Field Results" in text


def test_bulk_excel_export_has_file_centric_overview_columns():
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(to_excel_bytes_bulk(_sample_results(), "Lube Oil List Audit", "File")))
    sheet = workbook["Overview"]

    headers = [sheet.cell(row=4, column=col).value for col in range(1, 7)]
    assert headers == [
        "File",
        "Recommended Lubricant",
        "Recommended Grade",
        "UOM",
        "Overall Status",
        "Review Fields",
    ]
    assert sheet.cell(row=5, column=1).value == "MA-EPE255-F000643A/B"
    assert sheet.cell(row=5, column=5).value == "MATCHED"


def test_bulk_pdf_export_valid():
    data = to_pdf_bytes_bulk(_sample_results(), "Lube Oil List Audit")
    assert data[:4] == b"%PDF"


def test_bulk_csv_export_contains_all_equipment():
    text = to_csv_bytes_bulk(_sample_results(), "Lube Oil List Audit").decode("utf-8")
    assert "MA-EPE255-F000643A/B" in text
    assert "MF-EPE255-A112101" in text
    assert "MP-EPE255-P125003" in text


def test_bulk_exports_handle_single_row():
    results = [{"identifier": "Solo", "rows": [
        {"field": "X", "unit": "", "extracted_value": "1", "status": "MATCHED"}
    ], "summary": {"MATCHED": 1, "MISSING": 0, "MISMATCHED": 0, "OUT OF RANGE": 0, "TOTAL": 1}}]
    assert to_excel_bytes_bulk(results, "T")[:2] == b"PK"
    assert to_pdf_bytes_bulk(results, "T")[:4] == b"%PDF"
